"""docx。対象: 本文段落・表・ヘッダーフッター（設計書 4.6）。"""
from __future__ import annotations

from pathlib import Path

import docx

from ..core.models import Fragment
from .base import FileHandler


def _set_paragraph_text(paragraph, text: str) -> None:
    """段落テキストを差し替える。書式は先頭runのものに揃う。"""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _hf_location(kind: str, si: int, variant: str | None, pi: int) -> str:
    # 基本(primary)は既存呼び出し元との後方互換のためlocation文字列を変更しない。
    # 先頭ページ(first)/偶数ページ(even)はvariantタグを付与して区別する。
    prefix = f"{variant}:" if variant else ""
    return f"{kind}:{si}:{prefix}{pi}"


class DocxHandler(FileHandler):
    extensions = [".docx"]

    def _iter_hf_variants(self, doc, section):
        """(variantタグ, 有効フラグ, headerプロキシ, footerプロキシ) を列挙する。

        Word文書のヘッダー/フッターには、通常(primary)に加えて
        「先頭ページのみ別指定」時の first_page_header/footer、文書全体で
        「奇数/偶数ページ別指定」を有効にした場合の even_page_header/footer が
        存在しうる。これらのプロパティ自体を参照するだけではXMLへの副作用は
        ない（単なるプロキシオブジェクトの生成）。実際にpartが新規作成される
        のは、明示定義が存在しない状態で .paragraphs 等コンテンツへアクセス
        した場合のみ（_iter_locations側でガードする）。
        """
        yield None, True, section.header, section.footer
        yield ("first", section.different_first_page_header_footer,
               section.first_page_header, section.first_page_footer)
        yield ("even", doc.settings.odd_and_even_pages_header_footer,
               section.even_page_header, section.even_page_footer)

    def _iter_locations(self, doc):
        for i, p in enumerate(doc.paragraphs):
            yield f"para:{i}", p
        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, p in enumerate(cell.paragraphs):
                        loc = f"table:{ti}:{ri}:{ci}" if pi == 0 else f"table:{ti}:{ri}:{ci}:p{pi}"
                        yield loc, p
        for si, section in enumerate(doc.sections):
            for variant, enabled, header, footer in self._iter_hf_variants(doc, section):
                if not enabled:
                    continue
                # is_linked_to_previous が False の場合のみ、このセクションに
                # 明示的なヘッダー/フッター定義が存在する。python-docxでは
                # 明示定義が無い状態で .paragraphs にアクセスすると、新規の
                # 空パート(word/header1.xml等)が生成されてしまう副作用がある
                # ため、明示定義が確認できたものにしか触れない。
                if not header.is_linked_to_previous:
                    for pi, p in enumerate(header.paragraphs):
                        yield _hf_location("header", si, variant, pi), p
                if not footer.is_linked_to_previous:
                    for pi, p in enumerate(footer.paragraphs):
                        yield _hf_location("footer", si, variant, pi), p

    def read_fragments(self, path: Path) -> list[Fragment]:
        doc = docx.Document(str(path))
        return [Fragment(text=p.text, location=loc)
                for loc, p in self._iter_locations(doc) if p.text]

    def write_fragments(self, src: Path, dst: Path, masked: list[Fragment]) -> None:
        doc = docx.Document(str(src))
        by_loc = {f.location: f.text for f in masked}
        for loc, p in self._iter_locations(doc):
            if loc in by_loc and p.text != by_loc[loc]:
                _set_paragraph_text(p, by_loc[loc])
        doc.save(str(dst))
