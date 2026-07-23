"""pptx。対象: シェイプテキスト・表・スピーカーノート（設計書 4.6）。"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..core.models import Fragment
from .base import FileHandler


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def _has_txbody(shape) -> bool:
    """`.text_frame` に一度も触れずに `<p:txBody>` の有無を判定する。

    `Shape.has_text_frame` はテキストを入力したことがない装飾用シェイプ（アイ
    コン、コネクタ扱いされない図形等）に対しても常に |True| を返す。その状態で
    `.text_frame` へアクセスすると内部で `CT_Shape.get_or_add_txBody()` が呼ば
    れ、`<p:txBody>` が存在しない場合は新規に（空の段落を持つ）要素が追加され
    てしまう。これはマスキング処理が本来触れるべきでないシェイプへの構造的な
    副作用となるため、`element.txBody`（ZeroOrOne、読み取り専用で副作用のない
    プロパティ）で存在確認したものにしか `.text_frame` を使わない。
    """
    return shape.element.txBody is not None


class PptxHandler(FileHandler):
    extensions = [".pptx"]

    def _iter_shape_locations(self, shapes, prefix, in_group=False):
        """`shapes` 配下を再帰的に辿り (location, paragraph) を列挙する。

        `prefix` はこの呼び出しまでの階層を表す文字列（トップレベルは
        "slide:<s>"）。`MSO_SHAPE_TYPE.GROUP` のシェイプに遭遇した場合は
        `shape.shapes`（グループメンバーのコレクション）へ再帰し、以降の階層
        はすべて `:group:<gi>` セグメントを重ねて表現する（グループがグループ
        を含む任意の深さのネストに対応）。トップレベルのフォーマット
        （"slide:<s>:shape:<i>:para:<p>" / "slide:<s>:table:<i>:<r>:<c>:
        para:<p>"）は既存呼び出し元との後方互換のため変更しない。同一の再帰
        関数でトップレベル・ネストの両方を処理し、専用コードパスは持たない。
        """
        for idx, shape in enumerate(shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                child_prefix = (
                    f"{prefix}:group:{idx}" if in_group else f"{prefix}:shape:{idx}"
                )
                yield from self._iter_shape_locations(
                    shape.shapes, child_prefix, in_group=True)
                continue
            if shape.has_text_frame and _has_txbody(shape):
                tag = f"{prefix}:group:{idx}" if in_group else f"{prefix}:shape:{idx}"
                for pi, para in enumerate(shape.text_frame.paragraphs):
                    yield f"{tag}:para:{pi}", para
            if shape.has_table:
                tag = (
                    f"{prefix}:group:{idx}:table:{idx}"
                    if in_group else f"{prefix}:table:{idx}"
                )
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, para in enumerate(cell.text_frame.paragraphs):
                            yield f"{tag}:{ri}:{ci}:para:{pi}", para

    def _iter_locations(self, prs):
        for si, slide in enumerate(prs.slides):
            yield from self._iter_shape_locations(slide.shapes, f"slide:{si}")
            if slide.has_notes_slide:
                # has_notes_slideがTrueでも、ノートプレースホルダーがノート
                # スライドから削除されている場合はnotes_text_frameがNoneに
                # なりうる（python-pptx公式ドキュメントに明記されている挙動）。
                # Noneのまま.paragraphsへアクセスするとAttributeErrorで処理
                # 全体が中断するため、ここでガードする。
                notes_tf = slide.notes_slide.notes_text_frame
                if notes_tf is not None:
                    for pi, para in enumerate(notes_tf.paragraphs):
                        yield f"slide:{si}:notes:para:{pi}", para

    @staticmethod
    def _para_text(para) -> str:
        return "".join(run.text for run in para.runs)

    def read_fragments(self, path: Path) -> list[Fragment]:
        prs = Presentation(str(path))
        return [Fragment(text=self._para_text(p), location=loc)
                for loc, p in self._iter_locations(prs) if self._para_text(p)]

    def write_fragments(self, src: Path, dst: Path, masked: list[Fragment]) -> None:
        prs = Presentation(str(src))
        by_loc = {f.location: f.text for f in masked}
        for loc, para in self._iter_locations(prs):
            if loc in by_loc and self._para_text(para) != by_loc[loc]:
                _set_paragraph_text(para, by_loc[loc])
        prs.save(str(dst))
