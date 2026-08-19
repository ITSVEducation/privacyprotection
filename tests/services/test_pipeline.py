import re

import docx
import openpyxl
import pytest
from pptx import Presentation
from pptx.util import Inches
from privacyprotection.config import AppConfig
from privacyprotection.core.detector import Detector
from privacyprotection.core.dictionary import DictionaryDetector
from privacyprotection.core.mapping_io import read_mapping
from privacyprotection.core.masker import Masker
from privacyprotection.core.models import TOKEN_RE, Detection, Fragment
from privacyprotection.core.patterns import PatternDetector
from privacyprotection.handlers.docx_handler import DocxHandler
from privacyprotection.handlers.pptx_handler import PptxHandler
from privacyprotection.handlers.text_handler import TextHandler
from privacyprotection.handlers.xlsx_handler import XlsxHandler
from privacyprotection.services.pipeline import Pipeline


def make_pipeline(mode="token"):
    det = Detector([PatternDetector(),
                    DictionaryDetector({"山田太郎": "PERSON", "佐藤花子": "PERSON"})])
    return Pipeline(detector=det, mode=mode)


def test_mask_file_outputs_masked_and_mapping(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎 03-1234-5678", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    out, report = pl.mask_file(src, frags, dets)
    assert out.name == "memo_masked.txt"
    assert out.read_text(encoding="utf-8") == "【人名_1】 【電話_1】"
    assert src.read_text(encoding="utf-8") == "山田太郎 03-1234-5678"  # 元ファイル不変
    mapping = read_mapping(tmp_path / "memo_masked.txt.pmap.csv")
    assert len(mapping.entries) == 2
    assert report.category_counts == {"PERSON": 1, "PHONE": 1}


def test_existing_output_gets_numbered_suffix(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    (tmp_path / "memo_masked.txt").write_text("既存", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    out, _ = pl.mask_file(src, frags, dets)
    assert out.name == "memo_masked(2).txt"
    assert (tmp_path / "memo_masked.txt").read_text(encoding="utf-8") == "既存"


def test_redact_mode_produces_no_mapping(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline(mode="redact")
    frags, dets = pl.analyze_file(src)
    out, _ = pl.mask_file(src, frags, dets)
    assert out.read_text(encoding="utf-8") == "●●●●"
    assert not (tmp_path / "memo_masked.txt.pmap.csv").exists()


def test_disabled_detection_respected(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎と佐藤花子", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    for d in dets[0]:
        if d.text == "佐藤花子":
            d.enabled = False
    out, report = pl.mask_file(src, frags, dets)
    assert out.read_text(encoding="utf-8") == "【人名_1】と佐藤花子"


def test_mask_file_category_counts_reflect_post_overlap_resolution(tmp_path):
    """レビュー指摘（Finding 1）の回帰テスト: 手動追加(source="manual")の
    Detection が既存の自動検出と重なる場合、FileReport.category_counts は
    resolve_overlaps() 適用後に実際にマスクされた項目を反映しなければならない。

    渡された detections（解決前の生のリスト）をそのまま数えると、自動検出
    "cdefgh"(PERSON, 2-8) に完全に包含される手動追加 "de"(EMAIL, 3-5) が
    実際には（設計書4.2「完全に覆われる場合のみ丸ごと破棄」）マスク結果に
    一切現れないにもかかわらず、{"PERSON": 1, "EMAIL": 1} という存在しない
    EMAIL 項目まで水増しして数えてしまう。
    """
    text = "abcdefghij"
    src = tmp_path / "memo.txt"
    src.write_text(text, encoding="utf-8")
    pl = make_pipeline()
    frags = [Fragment(text=text, location="line:1")]

    auto = Detection(text="cdefgh", category="PERSON", start=2, end=8, source="ner")
    manual = Detection(text="de", category="EMAIL", start=3, end=5, source="manual")
    dets = [[auto, manual]]

    out, report = pl.mask_file(src, frags, dets)

    # マスク結果は PERSON の1トークンのみ（EMAIL は完全に包含されるため
    # 丸ごと破棄され、二重マスクも検出漏れも起きない）。
    assert out.read_text(encoding="utf-8") == "ab【人名_1】ij"
    # レポートの category_counts もそれと一致しなければならない
    # （解決前の生の件数を数えた場合の {"PERSON": 1, "EMAIL": 1} は誤り）。
    assert report.category_counts == {"PERSON": 1}


def test_mask_folder_shared_tokens_and_report(tmp_path):
    (tmp_path / "a.txt").write_text("山田太郎", encoding="utf-8")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.txt").write_text("山田太郎の件", encoding="utf-8")
    (tmp_path / "c.pdf").write_bytes(b"x")
    pl = make_pipeline()
    batch, mapping_path = pl.mask_folder(tmp_path)
    # 同一人物は全ファイルで同一トークン
    assert (tmp_path / "a_masked.txt").read_text(encoding="utf-8") == "【人名_1】"
    assert (tmp_path / "sub" / "b_masked.txt").read_text(encoding="utf-8") == "【人名_1】の件"
    assert mapping_path == tmp_path / "_folder.pmap.csv"
    assert (tmp_path / "_report.txt").exists()
    assert any(p.name == "c.pdf" for p, _ in batch.skipped)


def test_mask_folder_updates_mapping_before_each_masked_file(tmp_path, monkeypatch):
    """グローバル制約「対応表はマスク済ファイルより先に確定する」は、フォルダ
    一括処理でも各ファイル単位で保たれなければならない(バッチ全体の最後に
    1回だけ書くのでは不十分:途中で処理が中断されると、その時点までに書かれた
    マスク済ファイルのトークンが対応表に一切載っていない状態になり得る)。

    ここでは、各ファイルのマスク済み出力が実際にディスクへ書かれる直前の
    時点で、_folder.pmap.csv が既にそのファイルの新規トークンを含んだ状態に
    更新されていることを検証する。
    """
    (tmp_path / "a.txt").write_text("山田太郎", encoding="utf-8")
    (tmp_path / "b.txt").write_text("佐藤花子", encoding="utf-8")

    original_write = TextHandler.write_fragments
    checked: list[tuple[str, bool]] = []

    def spy_write(self, src, dst, masked):
        mapping_file = tmp_path / "_folder.pmap.csv"
        content = mapping_file.read_text(encoding="utf-8-sig") if mapping_file.exists() else ""
        for frag in masked:
            for tok in TOKEN_RE.findall(frag.text):
                checked.append((tok, tok in content))
        return original_write(self, src, dst, masked)

    monkeypatch.setattr(TextHandler, "write_fragments", spy_write)

    pl = make_pipeline()
    pl.mask_folder(tmp_path)

    assert checked  # スパイが実際に何かを検証したことを確認（誤って空なら無意味）
    assert all(present for _, present in checked), checked


def test_restore_file_roundtrip(tmp_path):
    src = tmp_path / "memo.txt"
    original = "山田太郎 03-1234-5678 佐藤花子"
    src.write_text(original, encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    restored_path, result = pl.restore_file(masked_path)  # pmap自動検出
    assert restored_path.read_text(encoding="utf-8") == original
    assert result.unknown_tokens == []


def test_restore_text_with_unknown_token(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    r = pl.restore_text("【人名_1】と【人名_9】",
                        masked_path.parent / (masked_path.name + ".pmap.csv"))
    assert r.text == "山田太郎と【人名_9】"
    assert r.unknown_tokens == ["【人名_9】"]


def test_broken_file_reported_not_raised(tmp_path):
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a zip")
    pl = make_pipeline()
    batch, _ = pl.mask_folder(tmp_path)
    assert len(batch.files) == 1
    assert batch.files[0].error is not None


def test_restore_file_numbered_masked_suffix(tmp_path):
    """マスク済み出力が連番付き（memo_masked(2).txt）の場合でも、その連番が
    復元ファイル名の語幹に紛れ込まない（stem.replace()によるバグの回帰テスト）。"""
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    # 1つ目の "memo_masked.txt" を先に存在させておき、2つ目を連番にする
    (tmp_path / "memo_masked.txt").write_text("既存", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    assert masked_path.name == "memo_masked(2).txt"

    restored_path, result = pl.restore_file(masked_path)
    assert restored_path.name == "memo_restored.txt"
    assert restored_path.read_text(encoding="utf-8") == "山田太郎"
    assert result.unknown_tokens == []


def test_restore_file_preserves_unrelated_masked_substring(tmp_path):
    """ファイル名にたまたま '_masked' を含む(本アプリ由来でない)場合、
    その部分文字列を誤って削ってはならない(サフィックス限定の除去であること)。"""
    src = tmp_path / "already_masked_by_someone_else.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    assert masked_path.name == "already_masked_by_someone_else_masked.txt"

    restored_path, _ = pl.restore_file(masked_path)
    # 語幹全体から "_masked" というサフィックスだけを取り除く。
    # 語中の "_masked" を盲目的に削ると
    # "already_by_someone_else_restored.txt" のような誤ったファイル名になる。
    assert restored_path.name == "already_masked_by_someone_else_restored.txt"


# Tests for Pipeline.from_config() and Pipeline.mask_text()

def test_from_config_with_custom_dictionary_token_mode():
    """from_config ファクトリが AppConfig.custom_dictionary を実際に
    DictionaryDetector へ配線することを確認する。

    検出対象には "PRJ-9981" という、NER（人名/組織/地名）にも定型パターン
    （メール・電話・郵便番号等）にも該当しない、辞書がなければ一切検出され
    得ない文字列を使う（"テスト太郎" のように NER 単体でも人名として拾われて
    しまう語句だと、from_config がカスタム辞書の配線をまるごと落としていても
    テストが偶然通ってしまうため不適切）。"""
    cfg = AppConfig(
        custom_dictionary={"PRJ-9981": "CUSTOM"},
        mask_mode="token",
        enabled_categories={"PERSON", "PHONE", "EMAIL", "CUSTOM"}
    )
    pl = Pipeline.from_config(cfg)

    text = "PRJ-9981の進捗を教えてください"
    masked, table, count = pl.mask_text(text)

    # カスタム辞書由来の検出がマスクされることを確認
    assert "PRJ-9981" not in masked
    assert count == 1
    assert len(table.entries) == 1
    assert table.entries[0].original == "PRJ-9981"
    assert table.entries[0].category == "CUSTOM"

    # 陰性対照: カスタム辞書エントリを持たない Pipeline では同じ文字列は
    # 検出されない。これにより、上の検出結果が本当に custom_dictionary の
    # 配線に起因すること（NER/パターン検出が偶然拾ったのではないこと）を
    # 確認する。
    cfg_no_dict = AppConfig(
        custom_dictionary={},
        mask_mode="token",
        enabled_categories={"PERSON", "PHONE", "EMAIL", "CUSTOM"}
    )
    pl_no_dict = Pipeline.from_config(cfg_no_dict)
    masked_no_dict, table_no_dict, count_no_dict = pl_no_dict.mask_text(text)
    assert "PRJ-9981" in masked_no_dict
    assert count_no_dict == 0
    assert len(table_no_dict.entries) == 0


def test_mask_text_token_mode_with_pii():
    """mask_text が token mode で複数の PII を検出・マスクし、正しい
    マスク済みテキスト、対応表、検出件数を返すことを確認する。"""
    pl = make_pipeline(mode="token")
    text = "山田太郎 03-1234-5678"

    masked, table, count = pl.mask_text(text)

    # 検出件数が正確に2件（人名1件、電話1件）
    assert count == 2
    # マスク済みテキストが生成されている
    assert "山田太郎" not in masked
    assert "03-1234-5678" not in masked
    # token mode では対応表エントリが存在する
    assert len(table.entries) > 0
    # マスク済みテキストに token が含まれる
    assert "【人名_" in masked
    assert "【電話_" in masked


def test_mask_text_token_mode_with_mapping_path(tmp_path):
    """mask_text が token mode で mapping_path を指定された場合、
    対応表エントリが存在するときのみ CSV ファイルを書き出すことを確認する。"""
    pl = make_pipeline(mode="token")
    text = "山田太郎"
    mapping_path = tmp_path / "test.pmap.csv"

    masked, table, count = pl.mask_text(text, mapping_path=mapping_path)

    # 対応表が書き出されていることを確認
    assert mapping_path.exists()
    assert count == 1
    assert len(table.entries) == 1
    # ファイルが読める形式であることを確認
    loaded = read_mapping(mapping_path)
    assert len(loaded.entries) == 1


def test_mask_text_redact_mode_regression_test(tmp_path):
    """mask_text が redact mode でも正しく検出件数を報告することを確認する。
    これは設計上、redact mode では対応表エントリが常に空になるため、
    「対応表が空 = 検出なし」という間違った推論をしてはならないことを
    保証する回帰テスト。実装では検出件数を別途 active_count で返すことで
    この区別が成立する。"""
    pl = make_pipeline(mode="redact")
    text = "山田太郎 03-1234-5678"

    masked, table, count = pl.mask_text(text)

    # redact mode ではテキストが ●●● で置き換わる
    assert "山田太郎" not in masked
    assert "03-1234-5678" not in masked
    assert "●●●●" in masked

    # 対応表エントリは ALWAYS 空（redact mode の仕様）
    assert len(table.entries) == 0

    # しかし検出件数は正確に報告される（非ゼロ）
    # これが非回帰である理由：table.entries の有無から count を
    # 推論していると、ここで count == 0 という誤った値になってしまう
    assert count == 2


def test_mask_text_no_pii():
    """mask_text が PII を含まないテキストを処理するとき、
    マスク済みテキストは元のまま、検出件数は 0 になることを確認する。"""
    pl = make_pipeline(mode="token")
    text = "これはPIIを含まない普通のテキストです"

    masked, table, count = pl.mask_text(text)

    # テキストが変わらない
    assert masked == text
    # 検出件数が 0
    assert count == 0
    # token mode でも対応表エントリは 0
    assert len(table.entries) == 0


def test_mask_text_redact_mode_no_pii_no_output_file(tmp_path):
    """mask_text が redact mode で PII が全くない場合、
    mapping_path を指定していてもファイルが作成されないことを確認する。
    （redact mode はそもそも対応表を作成しないモードなので、
    検出0件の場合に空の対応表ファイルを作る意味はない）"""
    pl = make_pipeline(mode="redact")
    text = "PII を含まないテキスト"
    mapping_path = tmp_path / "test.pmap.csv"

    masked, table, count = pl.mask_text(text, mapping_path=mapping_path)

    assert count == 0
    # redact mode ではそもそもファイルが作成されない
    assert not mapping_path.exists()


def test_mask_text_with_custom_masker():
    """mask_text が masker 引数を受け取ったとき、内部で新しい Masker を
    作り直さずそれをそのまま使う（＝呼び出しをまたいでカウンタ状態が
    共有される）ことを確認する。

    2回の呼び出しで「同じ」人名をマスクするテストだと、mask_text が渡された
    masker を無視して呼び出しごとに新しい Masker(mode="token") を作っていても
    両方【人名_1】になり区別が付かない。そこで2回の呼び出しで「異なる」人名を
    マスクし、採番が 1 → 2 と連続することを確認する。これは shared_masker が
    実際に使い回されていなければ成立しない（使い回されなければ2回目も
    【人名_1】から始まってしまう）。"""
    pl = make_pipeline(mode="token")
    shared_masker = Masker(mode="token")

    text1 = "山田太郎が来ました"
    text2 = "佐藤花子に聞いてください"

    masked1, _, _ = pl.mask_text(text1, masker=shared_masker)
    masked2, _, _ = pl.mask_text(text2, masker=shared_masker)

    tokens1 = re.findall(r"【人名_\d+】", masked1)
    tokens2 = re.findall(r"【人名_\d+】", masked2)

    # 1回目は_1、2回目は_2から始まり、カウンタが呼び出しをまたいで
    # 引き継がれていることを確認する。
    assert tokens1 == ["【人名_1】"]
    assert tokens2 == ["【人名_2】"]


# --- analyze_text / mask_text(detections=...) : クリップボード・プレビュー用 ---
#
# クリップボードのマスクにもファイルと同じ PreviewDialog を挟むため、生テキストを
# (Fragment, 検出リスト) に分解する analyze_text と、確認・編集後の検出リストを
# そのままマスクする mask_text(detections=...) を追加した。


def test_analyze_text_returns_clipboard_fragment_and_detections():
    """analyze_text が生テキストを location="clipboard" の Fragment 1つと、
    その検出リストに分解して返す。既知の PII を含めば検出は非空になる。"""
    pl = make_pipeline()
    frag, dets = pl.analyze_text("山田太郎 03-1234-5678")
    assert isinstance(frag, Fragment)
    assert frag.location == "clipboard"
    assert frag.text == "山田太郎 03-1234-5678"
    cats = {d.category for d in dets}
    assert "PERSON" in cats and "PHONE" in cats


def test_mask_text_uses_supplied_detections_without_redetecting():
    """detections を渡すと、内部で再検出せずそのリストだけをマスクする。
    ここでは実在の PII を含む文章に対し、手動で作った1件だけの検出リスト
    （文章の一部だけを指す）を渡し、その1件だけがマスクされ、本来なら自動
    検出されるはずの他の PII（電話番号）は素通りすることを確認する
    （＝渡したリストが尊重され、再検出が起きていないことの証明）。"""
    pl = make_pipeline(mode="token")
    text = "山田太郎 03-1234-5678"
    only = [Detection(text="山田太郎", category="PERSON",
                      start=0, end=4, source="manual")]
    masked, table, count = pl.mask_text(text, detections=only)
    assert count == 1
    assert "山田太郎" not in masked
    assert "03-1234-5678" in masked  # 再検出されていない証拠
    assert len(table.entries) == 1


def test_mask_text_with_all_disabled_detections_masks_nothing():
    """無効化された検出だけを渡すと、検出0件相当でマスクは何も起きない。"""
    pl = make_pipeline(mode="token")
    text = "山田太郎 03-1234-5678"
    _, dets = pl.analyze_text(text)
    for d in dets:
        d.enabled = False
    masked, table, count = pl.mask_text(text, detections=dets)
    assert count == 0
    assert masked == text
    assert len(table.entries) == 0


def test_mask_text_detections_none_falls_back_to_internal_detection():
    """detections=None（既定）のときは従来どおり内部検出でマスクする
    （後方互換）。"""
    pl = make_pipeline(mode="token")
    text = "山田太郎 03-1234-5678"
    masked, table, count = pl.mask_text(text)  # detections 未指定
    assert count == 2
    assert "山田太郎" not in masked
    assert "03-1234-5678" not in masked


# --- 再レビュー指摘: 形式別の「対象外」注記が設計書4.6の全項目を漏れなく
#     開示すること（44bdeca時点までの回帰の修正）---
#
# 以前は全Office形式に対して一律
# 「図形/テキストボックス内の文字と埋込オブジェクトは対象外です」という
# 注記を付けていたが、これは実態と食い違っていた:
#   - .pptx はシェイプ内テキスト（グループ化されたシェイプ含む）を実際に
#     マスクしており、注記は虚偽だった。
#   - .docx は脚注・文末脚注・コメントを対象外としているのに、注記はそれに
#     一切触れていなかった。
#   - .xlsx は数式内の文字列リテラル・定義名を対象外としているのに、注記は
#     それに一切触れていなかった。
# その直後の修正（前回ラウンド）でこれらの新事実は注記に追加されたが、今度は
# 設計書4.6の表がもともと要求していたテキストボックス/図形内文字・埋込画像・
# 埋込オブジェクト等の開示を丸ごと落としてしまう回帰を起こした（再レビュー
# 指摘）。以下は、新旧どちらの開示も欠落なく含む注記が各形式のFileReportに
# 付く（.pptxはシェイプテキストの虚偽claimなしに、埋込画像/埋込オブジェクト/
# スライドマスターのみ付く）ことを固定する回帰テスト。

def test_docx_report_notes_full_46_exclusion_list(tmp_path):
    d = docx.Document()
    d.add_paragraph("担当: 山田太郎")
    src = tmp_path / "memo.docx"
    d.save(src)
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    _, report = pl.mask_file(src, frags, dets)
    assert report.notes == [
        "テキストボックス/図形内文字・埋込画像・埋込オブジェクト・"
        "変更履歴の削除済みテキスト・脚注/文末脚注/コメントは対象外です"
    ]


def test_xlsx_report_notes_full_46_exclusion_list(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "山田太郎"
    src = tmp_path / "book.xlsx"
    wb.save(src)
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    _, report = pl.mask_file(src, frags, dets)
    assert report.notes == [
        "図形/テキストボックス内文字・埋込画像・埋込オブジェクト・"
        "ピボットキャッシュ・数式内の文字列リテラル・定義名は対象外です"
    ]


def test_pptx_report_notes_embedded_object_and_slide_master_only(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "山田太郎の報告"
    src = tmp_path / "deck.pptx"
    prs.save(src)
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    _, report = pl.mask_file(src, frags, dets)
    # シェイプテキストは実際にPptxHandlerがマスクしているため、それについての
    # 「対象外」claimは付かない。一方で埋込画像・埋込オブジェクト・
    # スライドマスターは設計書4.6の表どおり本当に対象外なので、注記は付く。
    assert report.notes == ["埋込画像・埋込オブジェクト・スライドマスターは対象外です"]


def test_plain_text_report_has_no_office_specific_note(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    _, report = pl.mask_file(src, frags, dets)
    assert report.notes == []


# --- マスク→復元ラウンドトリップ（Office形式）---
#
# 設計書8章は「全対応形式」でのラウンドトリップ検証を求めているが、これまで
# Pipeline経由のend-to-endテストはテキスト形式のみだった。各Office形式について
# 実ファイルを作り、mask_file → restore_file を通して、復元後の内容が
# 同じHandlerのread_fragments経由で元と完全に一致することを確認する。

def test_mask_restore_roundtrip_xlsx(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名簿"
    ws["A1"] = "山田太郎"
    ws["A2"] = "連絡先: 03-1234-5678"
    ws["A3"] = "佐藤花子"
    src = tmp_path / "book.xlsx"
    wb.save(src)

    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    assert masked_path != src
    restored_path, result = pl.restore_file(masked_path)

    handler = XlsxHandler()
    original = {f.location: f.text for f in handler.read_fragments(src)}
    restored = {f.location: f.text for f in handler.read_fragments(restored_path)}
    assert restored == original
    assert result.unknown_tokens == []


def test_mask_restore_roundtrip_docx(tmp_path):
    d = docx.Document()
    d.add_paragraph("担当: 山田太郎")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "佐藤花子"
    t.cell(0, 1).text = "090-1111-2222"
    d.sections[0].header.paragraphs[0].text = "社外秘 山田太郎"
    src = tmp_path / "doc.docx"
    d.save(src)

    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    restored_path, result = pl.restore_file(masked_path)

    handler = DocxHandler()
    original = {f.location: f.text for f in handler.read_fragments(src)}
    restored = {f.location: f.text for f in handler.read_fragments(restored_path)}
    assert restored == original
    assert result.unknown_tokens == []


def test_mask_restore_roundtrip_pptx(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "山田太郎の報告"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "連絡先: 090-1111-2222"
    slide.notes_slide.notes_text_frame.text = "佐藤花子に確認"
    src = tmp_path / "deck.pptx"
    prs.save(src)

    pl = make_pipeline()
    frags, dets = pl.analyze_file(src)
    masked_path, _ = pl.mask_file(src, frags, dets)
    restored_path, result = pl.restore_file(masked_path)

    handler = PptxHandler()
    original = {f.location: f.text for f in handler.read_fragments(src)}
    restored = {f.location: f.text for f in handler.read_fragments(restored_path)}
    assert restored == original
    assert result.unknown_tokens == []


def test_read_fragments_reads_without_detection(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline()
    frags = pl.read_fragments(src)
    assert [f.text for f in frags] == ["山田太郎"]


def test_read_fragments_rejects_unsupported_extension(tmp_path):
    src = tmp_path / "memo.exe"
    src.write_text("x", encoding="utf-8")
    pl = make_pipeline()
    with pytest.raises(ValueError):
        pl.read_fragments(src)


def test_analyze_file_reuses_given_fragments(tmp_path):
    src = tmp_path / "memo.txt"
    src.write_text("山田太郎", encoding="utf-8")
    pl = make_pipeline()
    frags = pl.read_fragments(src)
    src.unlink()  # 再読み込みしていればここで失敗する
    frags2, dets = pl.analyze_file(src, fragments=frags)
    assert frags2 is frags
    assert dets[0][0].category == "PERSON"


def test_restore_folder_roundtrip(tmp_path):
    (tmp_path / "a.txt").write_text("山田太郎です", encoding="utf-8")
    (tmp_path / "b.txt").write_text("佐藤花子さん", encoding="utf-8")
    pl = make_pipeline()
    pl.mask_folder(tmp_path)

    restored, warnings = pl.restore_folder(tmp_path)
    assert restored == 2
    assert warnings == []
    assert (tmp_path / "a_restored.txt").read_text(encoding="utf-8") == "山田太郎です"
    assert (tmp_path / "b_restored.txt").read_text(encoding="utf-8") == "佐藤花子さん"


def test_restore_folder_reports_failures_without_values(tmp_path):
    # マスク済み命名だが対応表がどこにもない → 1件の警告（値は含まない）
    (tmp_path / "orphan_masked.txt").write_text("【人名_9】", encoding="utf-8")
    pl = make_pipeline()
    restored, warnings = pl.restore_folder(tmp_path)
    assert restored == 0
    assert len(warnings) == 1
    assert "orphan_masked.txt" in warnings[0]
    assert "【人名_9】" not in warnings[0]


def test_restore_folder_skips_non_masked_files(tmp_path):
    (tmp_path / "plain.txt").write_text("平文", encoding="utf-8")
    pl = make_pipeline()
    restored, warnings = pl.restore_folder(tmp_path)
    assert restored == 0
    assert warnings == []
    assert not (tmp_path / "plain_restored.txt").exists()
