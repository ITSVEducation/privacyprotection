from pptx import Presentation
from pptx.util import Inches
from privacyprotection.handlers.pptx_handler import PptxHandler

h = PptxHandler()

def make_pptx(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "山田太郎の報告"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "連絡先: 090-1111-2222"
    slide.notes_slide.notes_text_frame.text = "佐藤花子に確認"
    p = tmp_path / "test.pptx"
    prs.save(p)
    return p

def test_reads_shapes_and_notes(tmp_path):
    frags = h.read_fragments(make_pptx(tmp_path))
    all_text = " ".join(f.text for f in frags)
    assert "山田太郎の報告" in all_text
    assert "連絡先: 090-1111-2222" in all_text
    assert "佐藤花子に確認" in all_text

def test_write_replaces_text(tmp_path):
    src = make_pptx(tmp_path)
    frags = h.read_fragments(src)
    for f in frags:
        f.text = f.text.replace("山田太郎", "【人名_1】")
    dst = tmp_path / "out.pptx"
    h.write_fragments(src, dst, frags)
    out_frags = h.read_fragments(dst)
    all_text = " ".join(f.text for f in out_frags)
    assert "【人名_1】の報告" in all_text
    assert "山田太郎" not in all_text
