import openpyxl
from openpyxl.comments import Comment
from privacyprotection.handlers.xlsx_handler import XlsxHandler

h = XlsxHandler()

def make_book(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "名簿"
    ws["A1"] = "山田太郎"
    ws["B1"] = 42          # 数値 → 対象外
    ws["A2"] = "=SUM(1,2)" # 数式 → 対象外
    ws["A3"] = "佐藤花子"
    ws["A3"].comment = Comment("連絡先: 090-1111-2222", "author")
    ws.oddHeader.center.text = "社外秘 山田太郎"
    ws.oddFooter.center.text = "作成者 山田太郎"
    p = tmp_path / "test.xlsx"
    wb.save(p)
    return p

def test_reads_string_cells_comments_header_footer_sheetname(tmp_path):
    frags = h.read_fragments(make_book(tmp_path))
    texts = {f.location: f.text for f in frags}
    assert texts["名簿!A1"] == "山田太郎"
    assert texts["名簿!A3"] == "佐藤花子"
    assert texts["名簿!A3#comment"] == "連絡先: 090-1111-2222"
    assert texts["名簿#header"] == "社外秘 山田太郎"
    assert texts["名簿#footer"] == "作成者 山田太郎"
    assert texts["sheetname:名簿"] == "名簿"
    assert "名簿!B1" not in texts   # 数値は対象外
    assert "名簿!A2" not in texts   # 数式は対象外

def test_write_replaces_only_masked_text(tmp_path):
    src = make_book(tmp_path)
    frags = h.read_fragments(src)
    for f in frags:
        f.text = f.text.replace("山田太郎", "【人名_1】")
    dst = tmp_path / "out.xlsx"
    h.write_fragments(src, dst, frags)
    wb = openpyxl.load_workbook(dst)
    ws = wb["名簿"]
    assert ws["A1"].value == "【人名_1】"
    assert ws["B1"].value == 42
    assert ws["A3"].value == "佐藤花子"
    assert ws.oddHeader.center.text == "社外秘 【人名_1】"
    assert ws.oddFooter.center.text == "作成者 【人名_1】"


def test_write_replaces_comment_text(tmp_path):
    # コメント書き戻し分岐は、実際にテキストが変化するケースで検証する。
    src = make_book(tmp_path)
    frags = h.read_fragments(src)
    for f in frags:
        if f.location == "名簿!A3#comment":
            f.text = f.text.replace("090-1111-2222", "【電話_1】")
    dst = tmp_path / "out_comment.xlsx"
    h.write_fragments(src, dst, frags)
    wb = openpyxl.load_workbook(dst)
    assert wb["名簿"]["A3"].comment.text == "連絡先: 【電話_1】"


def test_write_renames_sheet(tmp_path):
    # シート名書き戻し分岐（"sheetname:..." location）の検証。
    # 全角括弧のトークンはExcelのシート名禁止文字(: \ / ? * [ ])に該当しない。
    src = make_book(tmp_path)
    frags = h.read_fragments(src)
    for f in frags:
        if f.location == "sheetname:名簿":
            f.text = "【人名_1】"
    dst = tmp_path / "out_rename.xlsx"
    h.write_fragments(src, dst, frags)
    wb = openpyxl.load_workbook(dst)
    assert wb.sheetnames == ["【人名_1】"]


def test_string_cell_with_equals_prefix_is_detected_via_data_type(tmp_path):
    # cell.value への通常代入は先頭"="を数式と自動判定してしまうため、
    # 「文字列セルだが内容がたまたま=で始まる」状態は内部属性を直接操作して再現する。
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "書式"
    ws["A1"] = "some placeholder"
    cell = ws["A1"]
    cell._value = "=見た目は数式だが実は文字列 090-1234-5678"
    cell.data_type = "s"
    ws["A2"] = "=SUM(1,2)"  # 本物の数式は引き続き対象外
    ws["B2"] = 123          # 数値も引き続き対象外
    p = tmp_path / "eq.xlsx"
    wb.save(p)

    # data_type はセルの t 属性としてXMLに保存されるため、内容とは独立して
    # ラウンドトリップすることを確認する。
    reloaded = openpyxl.load_workbook(p)
    assert reloaded["書式"]["A1"].data_type == "s"
    assert reloaded["書式"]["A2"].data_type == "f"
    assert reloaded["書式"]["B2"].data_type == "n"

    frags = h.read_fragments(p)
    texts = {f.location: f.text for f in frags}
    assert texts["書式!A1"] == "=見た目は数式だが実は文字列 090-1234-5678"
    assert "書式!A2" not in texts  # 本物の数式は対象外のまま
    assert "書式!B2" not in texts  # 数値は対象外のまま


def test_reads_left_and_even_header_and_first_footer_variants(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "帳票"
    ws.oddHeader.left.text = "作成: 山田太郎"
    ws.evenHeader.center.text = "偶数ページ 佐藤花子"
    ws.firstFooter.right.text = "初回のみ 090-0000-1111"
    p = tmp_path / "hf.xlsx"
    wb.save(p)

    frags = h.read_fragments(p)
    texts = {f.location: f.text for f in frags}
    assert texts["帳票#header:odd:left"] == "作成: 山田太郎"
    assert texts["帳票#header:even:center"] == "偶数ページ 佐藤花子"
    assert texts["帳票#footer:first:right"] == "初回のみ 090-0000-1111"


def test_write_replaces_header_variant_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "帳票"
    ws.oddHeader.left.text = "作成: 山田太郎"
    src = tmp_path / "hf_src.xlsx"
    wb.save(src)

    frags = h.read_fragments(src)
    for f in frags:
        if f.location == "帳票#header:odd:left":
            f.text = f.text.replace("山田太郎", "【人名_1】")
    dst = tmp_path / "hf_out.xlsx"
    h.write_fragments(src, dst, frags)

    reloaded = openpyxl.load_workbook(dst)
    assert reloaded["帳票"].oddHeader.left.text == "作成: 【人名_1】"


def test_write_preserves_string_data_type_on_equals_prefix_cell(tmp_path):
    # 数式と自動判定されないよう、data_type = "s" な文字列セルを作る。
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "データ"
    ws["A1"] = "placeholder"
    cell = ws["A1"]
    cell._value = "=電話番号 090-1234-5678"
    cell.data_type = "s"
    src = tmp_path / "eq_mask_src.xlsx"
    wb.save(src)

    # リード→マスク→ライト
    frags = h.read_fragments(src)
    for f in frags:
        if f.location == "データ!A1":
            # 電話番号部分のみをマスクし、先頭の "=" は残す
            f.text = f.text.replace("090-1234-5678", "【電話_1】")
    dst = tmp_path / "eq_mask_out.xlsx"
    h.write_fragments(src, dst, frags)

    # ラウンドトリップ後、data_type は "s" のまま、値は正しく更新されていることを確認
    reloaded = openpyxl.load_workbook(dst)
    cell_reloaded = reloaded["データ"]["A1"]
    assert cell_reloaded.data_type == "s"
    assert cell_reloaded.value == "=電話番号 【電話_1】"
