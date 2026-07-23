import docx
from privacyprotection.handlers.docx_handler import DocxHandler

h = DocxHandler()

def make_doc(tmp_path):
    d = docx.Document()
    d.add_paragraph("担当: 山田太郎")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "佐藤花子"
    t.cell(0, 1).text = "090-1111-2222"
    d.sections[0].header.paragraphs[0].text = "社外秘 山田太郎"
    p = tmp_path / "test.docx"
    d.save(p)
    return p

def test_reads_paragraphs_tables_header(tmp_path):
    frags = h.read_fragments(make_doc(tmp_path))
    texts = {f.location: f.text for f in frags}
    assert texts["para:0"] == "担当: 山田太郎"
    assert texts["table:0:0:0"] == "佐藤花子"
    assert texts["table:0:0:1"] == "090-1111-2222"
    assert any(loc.startswith("header:") and "山田太郎" in t
               for loc, t in texts.items())

def test_write_replaces_text(tmp_path):
    src = make_doc(tmp_path)
    frags = h.read_fragments(src)
    for f in frags:
        f.text = f.text.replace("山田太郎", "【人名_1】")
    dst = tmp_path / "out.docx"
    h.write_fragments(src, dst, frags)
    d = docx.Document(str(dst))
    assert d.paragraphs[0].text == "担当: 【人名_1】"
    assert d.tables[0].cell(0, 0).text == "佐藤花子"
    assert "【人名_1】" in d.sections[0].header.paragraphs[0].text
